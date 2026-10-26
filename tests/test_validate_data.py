"""Tests de ``src/futbol/db/validate_data.py`` contra una BD sintética.

No hacen ninguna llamada de red. Construyen escenarios mínimos
insertando directamente en la BD (no via ``load_data.py``, que ya tiene
sus propios tests) para aislar cada validación.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
import pytest

from futbol.db.database import get_connection, init_db
from futbol.db.validate_data import (
    _duplicados_matches,
    _fechas_fuera_de_rango,
    _fks_huerfanas,
    _null_percentages,
    _scores_negativos,
    run_validation,
)

REPO_SCHEMA_PATH = Path(__file__).resolve().parent.parent / "db" / "schema_v3.sql"


@pytest.fixture
def conn(tmp_path: Path):
    db_path = tmp_path / "futbol.db"
    init_db(db_path, REPO_SCHEMA_PATH)
    connection = get_connection(db_path)
    connection.execute("INSERT INTO leagues (code) VALUES ('TEST-League')")
    connection.execute(
        "INSERT INTO seasons (league_id, code) VALUES "
        "((SELECT league_id FROM leagues WHERE code = 'TEST-League'), '2024')"
    )
    connection.execute("INSERT INTO teams (canonical_name) VALUES ('Team A')")
    connection.execute("INSERT INTO teams (canonical_name) VALUES ('Team B')")
    connection.commit()
    yield connection
    connection.close()


def _season_id(conn) -> int:
    return conn.execute("SELECT season_id FROM seasons").fetchone()[0]


def _insert_match(conn, match_date: str, home_score=None, away_score=None) -> int:
    season_id = _season_id(conn)
    cur = conn.execute(
        "INSERT INTO matches (season_id, home_team_id, away_team_id, match_date, home_score, away_score) "
        "VALUES (?, 1, 2, ?, ?, ?)",
        (season_id, match_date, home_score, away_score),
    )
    conn.commit()
    return cur.lastrowid


# ---------------------------------------------------------------------------
# Validaciones individuales
# ---------------------------------------------------------------------------

def test_null_percentages_calcula_correctamente(conn):
    _insert_match(conn, "2024-08-16 19:00:00", home_score=1, away_score=0)
    _insert_match(conn, "2024-08-17 19:00:00", home_score=None, away_score=None)

    porcentajes = _null_percentages(conn, "matches")

    assert porcentajes["home_score"] == 50.0
    assert porcentajes["away_score"] == 50.0
    assert porcentajes["match_date"] == 0.0


def test_null_percentages_tabla_vacia_devuelve_vacio(conn):
    assert _null_percentages(conn, "matches") == {}


def test_duplicados_matches_detecta_llave_repetida(conn):
    season_id = _season_id(conn)
    conn.execute(
        "INSERT INTO matches (season_id, home_team_id, away_team_id, match_date) "
        "VALUES (?, 1, 2, '2024-08-16 19:00:00')",
        (season_id,),
    )
    conn.execute(
        "INSERT INTO matches (season_id, home_team_id, away_team_id, match_date) "
        "VALUES (?, 1, 2, '2024-08-16 19:00:00')",
        (season_id,),
    )
    conn.commit()

    duplicados = _duplicados_matches(conn)

    assert len(duplicados) == 1
    assert duplicados[0][4] == 2  # cantidad de filas repetidas


def test_duplicados_matches_sin_repetidos_devuelve_vacio(conn):
    _insert_match(conn, "2024-08-16 19:00:00")

    assert _duplicados_matches(conn) == []


def test_fechas_fuera_de_rango_detecta_año_invalido(conn):
    _insert_match(conn, "1800-01-01 12:00:00")
    _insert_match(conn, "2024-08-16 19:00:00")

    fuera_de_rango = _fechas_fuera_de_rango(conn)

    assert len(fuera_de_rango) == 1
    assert fuera_de_rango[0][1] == "1800-01-01 12:00:00"


def test_scores_negativos_detecta_marcador_invalido(conn):
    _insert_match(conn, "2024-08-16 19:00:00", home_score=-1, away_score=2)
    _insert_match(conn, "2024-08-17 19:00:00", home_score=1, away_score=0)

    negativos = _scores_negativos(conn)

    assert len(negativos) == 1
    assert negativos[0][1] == -1


def test_fks_huerfanas_vacio_en_bd_sana(conn):
    _insert_match(conn, "2024-08-16 19:00:00")

    assert _fks_huerfanas(conn) == []


# ---------------------------------------------------------------------------
# run_validation (reporte completo)
# ---------------------------------------------------------------------------

def test_run_validation_bd_sana_reporta_cero_errores(conn, tmp_path: Path):
    _insert_match(conn, "2024-08-16 19:00:00", home_score=1, away_score=0)

    path = run_validation(conn, report_dir=tmp_path, today=date(2026, 10, 14))

    assert path == tmp_path / "validation_2026-10-14.xlsx"
    wb = openpyxl.load_workbook(path)
    resumen = {row[0].value: row[1].value for row in wb["resumen"].iter_rows(min_row=2)}
    assert resumen["total_errores"] == 0
    assert resumen["total_matches"] == 1


def test_run_validation_detecta_y_reporta_los_3_tipos_de_error_en_matches(conn, tmp_path: Path):
    _insert_match(conn, "1800-01-01 12:00:00")  # fecha fuera de rango
    _insert_match(conn, "2024-08-16 19:00:00", home_score=-1, away_score=0)  # score negativo
    season_id = _season_id(conn)
    conn.execute(
        "INSERT INTO matches (season_id, home_team_id, away_team_id, match_date) "
        "VALUES (?, 1, 2, '2024-08-17 19:00:00')", (season_id,),
    )
    conn.execute(
        "INSERT INTO matches (season_id, home_team_id, away_team_id, match_date) "
        "VALUES (?, 1, 2, '2024-08-17 19:00:00')", (season_id,),
    )  # duplicado
    conn.commit()

    path = run_validation(conn, report_dir=tmp_path, today=date(2026, 10, 14))

    wb = openpyxl.load_workbook(path)
    resumen = {row[0].value: row[1].value for row in wb["resumen"].iter_rows(min_row=2)}
    assert resumen["fechas_fuera_de_rango"] == 1
    assert resumen["scores_negativos"] == 1
    assert resumen["matches_duplicados"] == 1
    assert resumen["total_errores"] == 3

    checks_reportados = {row[1].value for row in wb["errores"].iter_rows(min_row=2)}
    assert checks_reportados == {"fecha_fuera_de_rango", "score_negativo", "duplicado"}


def test_run_validation_crea_report_dir_si_no_existe(conn, tmp_path: Path):
    report_dir = tmp_path / "subcarpeta" / "reports"

    path = run_validation(conn, report_dir=report_dir, today=date(2026, 10, 14))

    assert path.exists()
    assert report_dir.is_dir()
